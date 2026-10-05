"""Synthetic class for trying vignettes: portraits (cover + grid) per student, group photos,
teachers, and a filled class workbook in demo_class/<template>/."""
import argparse
import random
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

from make_demo_data import make_photo
from vignette.config import class_photo_labels, field_slots, load_template_config
from vignette.table import CLASS_NAME, CLASS_SHEET, COMMON_SHEET, STUDENTS_SHEET, TEACHERS_SHEET, write_table
from vignette.cli_common import load_manifest

HERE = Path(__file__).resolve().parent
FIRST = ["Алия", "Эльдар", "Атай", "Айжан", "Салия", "Аяна", "Темир", "Диана", "Азамат", "Камила"]
LAST = ["Абишева", "Азимбеков", "Айнабеков", "Мамытова", "Алмазбекова", "Аманбаева", "Бейшекадыров",
        "Дононбаева", "Мамадалиев", "Шестопалова", "Константинопольская"]
QUOTES = [None, "Вперёд и только вперёд", "Stressed, depressed, but well dressed",
          "Лучшие годы\nмы провели вместе", "67"]
FIELD_VALUES = {"Год выпуска": "2025", "Класс": "11 “Б”", "Город": "Г.БИШКЕК",
                "Школа": "ШКОЛА-ГИМНАЗИЯ №13\nГОРОДА БИШКЕК"}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--template", required=True)
    ap.add_argument("--students", type=int, default=30)
    args = ap.parse_args()
    rnd = random.Random(7)
    manifest = load_manifest(args.template)
    cfg = load_template_config(args.template, manifest)
    root = HERE / "demo_class" / args.template
    shutil.rmtree(root, ignore_errors=True)
    for sub in ("портреты", "групповые", "учителя"):
        (root / "photos" / sub).mkdir(parents=True)

    names, seen = [], set()
    while len(names) < args.students:
        name = f"{rnd.choice(LAST)} {rnd.choice(FIRST)}"
        if name not in seen:
            seen.add(name)
            names.append(name)
    names.sort()
    for i, name in enumerate(names, 1):
        make_photo(root / "photos" / "портреты" / f"{i:02d}-обложка.jpg", f"{i} {name.split()[1]}", True, rnd)
        make_photo(root / "photos" / "портреты" / f"{i:02d}-сетка.jpg", f"{i}", True, rnd)
    groups = [f"общая{i:02d}" for i in range(1, 13)]
    for g in groups:
        make_photo(root / "photos" / "групповые" / f"{g}.jpg", g, False, rnd)
    for i in range(1, 15):
        make_photo(root / "photos" / "учителя" / f"учитель{i:02d}.jpg", f"У{i}", True, rnd)
    for label in class_photo_labels(cfg):
        make_photo(root / "photos" / "учителя" / f"{label}.jpg", label[:12], True, rnd)

    xlsx = root / "класс.xlsx"
    write_table(cfg, manifest, xlsx)
    wb = load_workbook(xlsx)
    ws = wb[CLASS_SHEET]
    labels = {f["label"] for f in field_slots(cfg, manifest)} | set(class_photo_labels(cfg))
    for row in ws.iter_rows(min_row=2):
        label = row[0].value
        if label == CLASS_NAME:
            row[1].value = "11 Б"
        elif label in FIELD_VALUES:
            row[1].value = FIELD_VALUES[label]
        elif label in class_photo_labels(cfg):
            row[1].value = label
        elif label in labels and label.startswith("Друзья"):
            row[1].value = rnd.choice(FIRST if "имя" in label else LAST)
    titles = [p["title"] for p in cfg["pages"] if not p.get("required")]
    for i, name in enumerate(names, 1):
        chosen = rnd.sample(groups, rnd.randint(0, 4))
        pages = None if i % 3 else ", ".join(rnd.sample(titles, min(2, len(titles))))
        wb[STUDENTS_SHEET].append([name, f"{i:02d}-обложка", f"{i:02d}-сетка", rnd.choice(QUOTES),
                                   ", ".join(chosen) or None, pages])
    if TEACHERS_SHEET in wb.sheetnames:
        for i in range(1, 15):
            wb[TEACHERS_SHEET].append([f"учитель{i:02d}", f"Учитель{i} Имя Отчество / предмет {i}"])
    for g in groups:
        wb[COMMON_SHEET].append([g])
    wb.save(xlsx)
    print(root)


if __name__ == "__main__":
    main()
