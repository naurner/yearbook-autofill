from openpyxl import load_workbook

from vhelpers import config, manifest
from vignette.table import (CLASS_SHEET, COMMON_SHEET, PAGES_SHEET, STUDENTS_HEADER, STUDENTS_SHEET,
                            TEACHERS_SHEET, read_table, write_table)


def make(tmp_path, cfg=None):
    path = tmp_path / "class.xlsx"
    write_table(cfg or config(), manifest(), path)
    return path


def test_write_table_sheets(tmp_path):
    wb = load_workbook(make(tmp_path))
    assert wb.sheetnames == ["Инструкция", CLASS_SHEET, STUDENTS_SHEET, TEACHERS_SHEET, COMMON_SHEET, PAGES_SHEET]
    labels = [r[0] for r in wb[CLASS_SHEET].iter_rows(min_row=2, values_only=True)]
    assert labels == ["Название класса", "Год", "Класс"]          # personal field hidden
    assert [c.value for c in wb[STUDENTS_SHEET][1]] == STUDENTS_HEADER
    pages = list(wb[PAGES_SHEET].iter_rows(min_row=2, values_only=True))
    assert pages[0][:3] == ("cover", "Обложка", "да") and pages[2][3] == 3


def test_no_teachers_sheet_without_teachers_page(tmp_path):
    cfg = config()
    cfg["pages"] = [p for p in cfg["pages"] if p["role"] != "teachers"]
    assert TEACHERS_SHEET not in load_workbook(make(tmp_path, cfg)).sheetnames


def fill(path, class_rows=(), students=(), teachers=(), common=()):
    wb = load_workbook(path)
    ws = wb[CLASS_SHEET]
    for label, value in class_rows:
        for row in ws.iter_rows(min_row=2):
            if row[0].value == label:
                row[1].value = value
    for r in students:
        wb[STUDENTS_SHEET].append(list(r))
    for r in teachers:
        wb[TEACHERS_SHEET].append(list(r))
    for r in common:
        wb[COMMON_SHEET].append([r])
    wb.save(path)


def test_read_table_roundtrip(tmp_path):
    path = make(tmp_path)
    fill(path, class_rows=[("Название класса", "11 А"), ("Год", 2025)],
         students=[("Абишева Алия", "c1.jpg", "p1", 67, "g1, g2;g3", "3, мои друзья"),
                   (None, None, None, None, None, None),
                   ("Азимбеков Эльдар", "c2", "p2", "Вперёд\nи вверх", None, None)],
         teachers=[("t1.jpg", "Иванова Анна / математика")], common=["g4", "g5"])
    cls = read_table(config(), manifest(), path)
    assert not cls["errors"]
    assert cls["class_name"] == "11 А"
    assert cls["fields"] == {"Год": "2025", "Класс": None}
    a, b = cls["students"]
    assert (a["surname"], a["first"], a["quote"], a["group"]) == ("Абишева", "Алия", "67", ["g1", "g2", "g3"])
    assert a["pages"] == ["03", "02"] and a["row"] == 2
    assert b["pages"] is None and b["quote"] == "Вперёд\nи вверх" and b["row"] == 4
    assert cls["teachers"] == [{"photo": "t1.jpg", "caption": "Иванова Анна\nматематика"}]
    assert cls["common_photos"] == ["g4", "g5"]


def test_read_table_errors_and_duplicates(tmp_path):
    path = make(tmp_path)
    fill(path, students=[("Абишева Алия", "c", "p", None, None, "99"),
                         (None, "c", "p", None, None, None),
                         ("Абишева Алия", "c", "p", None, None, None)])
    cls = read_table(config(), manifest(), path)
    assert any("99" in e and "строка 2" in e for e in cls["errors"])
    assert any("строка 3" in e and "имя" in e for e in cls["errors"])
    assert cls["students"][-1]["name"] == "Абишева Алия (2)"
    assert any("Абишева Алия" in w for w in cls["warnings"])
    assert cls["class_name"] == "class"                               # falls back to the file name


def test_fixed_fields_not_in_table(tmp_path):
    cfg = config()
    cfg["fields"].append({"label": "Заголовок", "layers": ["02:01"], "value": "НАШ КЛАСС"})
    labels = [r[0] for r in load_workbook(make(tmp_path, cfg))[CLASS_SHEET].iter_rows(min_row=2, values_only=True)]
    assert "Заголовок" not in labels


def test_no_extra_pages_token(tmp_path):
    path = make(tmp_path)
    fill(path, students=[("Абишева Алия", "c", "p", None, None, "нет"), ("Азимбеков Эльдар", "c", "p", None, None, "-")])
    cls = read_table(config(), manifest(), path)
    assert not cls["errors"] and [s["pages"] for s in cls["students"]] == [[], []]
