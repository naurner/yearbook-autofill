from openpyxl import load_workbook

from autofill.excel_io import FIELDS_SHEET, NO_SLOT_FILL, read_filled, write_template
from autofill.manifest import PHOTO_COL, build_manifest
from helpers import grid_layout

SHEET = "Стр 01 (4 фото)"
CAPTION = "Подпись: Имя Фамилия"


def make(tmp_path):
    m = build_manifest("T", [("01", grid_layout())])
    path = tmp_path / "t.xlsx"
    write_template(m, path)
    return m, path


def test_template_structure(tmp_path):
    _, path = make(tmp_path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Инструкция", FIELDS_SHEET, SHEET]
    ws = wb[SHEET]
    assert [c.value for c in ws[1]] == ["№", CAPTION, PHOTO_COL]
    assert [ws.cell(row=r, column=1).value for r in range(2, 6)] == [1, 2, 3, 4]
    assert [c.value for c in wb[FIELDS_SHEET][2]] == ["p01_01", "01", "П1", "правая половина, низ", "11 “А”", None]


def test_cells_without_slot_are_gray(tmp_path):
    m = build_manifest("T", [("01", grid_layout())])
    m["pages"][0]["groups"][0]["records"][0]["texts"] = []
    path = tmp_path / "t.xlsx"
    write_template(m, path)
    ws = load_workbook(path)[SHEET]
    assert ws["B2"].fill.fgColor.rgb.endswith(NO_SLOT_FILL.fgColor.rgb[-6:])
    assert ws["B3"].fill.fill_type is None


def test_roundtrip_filled_values(tmp_path):
    m, path = make(tmp_path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    ws["B2"], ws["C2"], ws["C3"] = "Мария\nИванова", "ivanova.jpg", "petrov"
    wb[FIELDS_SHEET]["F2"] = 11
    wb.save(path)
    data = read_filled(m, path)
    assert data["fields"] == {"p01_01": "11"}
    rows = data["groups"][SHEET]
    assert rows[0] == {CAPTION: "Мария\nИванова", PHOTO_COL: "ivanova.jpg"}
    assert rows[1] == {CAPTION: None, PHOTO_COL: "petrov"}
    assert len(rows) == 4 and data["errors"] == []


def test_extra_rows_are_errors(tmp_path):
    m, path = make(tmp_path)
    wb = load_workbook(path)
    wb[SHEET]["C7"] = "extra.jpg"
    wb.save(path)
    data = read_filled(m, path)
    assert len(data["errors"]) == 1 and "строка 7" in data["errors"][0]


def test_missing_sheet_is_warning(tmp_path):
    m, path = make(tmp_path)
    wb = load_workbook(path)
    del wb[SHEET]
    wb.save(path)
    data = read_filled(m, path)
    assert data["groups"][SHEET] == [] and data["warnings"]
