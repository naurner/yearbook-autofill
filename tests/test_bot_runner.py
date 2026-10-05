import fitz
import pytest
from openpyxl import load_workbook
from PIL import Image

from bot.flows import App
from bot.runner import run_action
from bot.store import Store

ADMIN = 1


@pytest.fixture
def world(tmp_path):
    for sub, nums in {"портретки": range(1001, 1005), "общие": range(2001, 2006)}.items():
        (tmp_path / "photos" / sub).mkdir(parents=True)
        for n in nums:
            Image.new("RGB", (60, 80), (n % 255, 80, 160)).save(tmp_path / "photos" / sub / f"IMG_{n}.jpg")
    app = App(store=Store(tmp_path / "db.sqlite"), admins={ADMIN}, data_dir=tmp_path / "data")
    code = app.store.create_class(template="Flight", title="11 Б", expected=2, tariff=1,
                                  source={"kind": "folder", "path": str(tmp_path / "photos")},
                                  fields={"Год выпуска": "2025"}, class_photos={}, common=["2001", "2002"], admin_id=1)
    for uid, name, pages in [(10, "Абишева Алия", ["03"]), (20, "Азимбеков Эльдар", [])]:
        st = app.store.join(code, uid)
        app.store.update_student(st["id"], name=name, cover=str(1000 + uid // 10), portrait="1003", quote="Вперёд",
                                 grp=["2003"], pages=pages, status="filled")
    app.store.join(code, 30)                                   # still filling
    return app, code


def test_previews_sent_to_ready_students(world):
    app, code = world
    outs = run_action(app, ("previews", code), ADMIN)
    docs = [o for o in outs if o.document]
    assert sorted(o.chat for o in docs) == [10, 20]
    assert all(o.buttons[0][0][1] == "approve" for o in docs)
    with fitz.open(next(o.document for o in docs if o.chat == 10)) as d:
        assert d.page_count == 5                     # cover, grid, «Наш класс», «Цитаты» + «Мои друзья»
    assert app.store.get_class(code)["status"] == "preview"
    assert {s["status"] for s in app.store.students(code, {"previewed"})} == {"previewed"}
    assert any(o.chat == ADMIN and "разосланы (2)" in o.text for o in outs)


def test_preview_one_and_missing_photo(world):
    app, code = world
    st = app.store.student(code, 20)
    outs = run_action(app, ("preview_one", code, st["id"]), ADMIN)
    assert [o.chat for o in outs if o.document] == [20]
    app.store.update_student(st["id"], cover="9999")
    outs = run_action(app, ("previews", code), ADMIN)
    assert not any(o.document for o in outs) and "9999" in outs[-1].text


def test_xlsx_and_remind(world):
    app, code = world
    outs = run_action(app, ("xlsx", code), ADMIN)
    wb = load_workbook(outs[0].document)
    rows = list(wb["Ученики"].iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["Абишева Алия", "Азимбеков Эльдар"] and rows[1][5] == "нет"
    outs = run_action(app, ("remind", code), ADMIN)
    assert [o.chat for o in outs] == [30, ADMIN]
