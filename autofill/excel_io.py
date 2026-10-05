"""Operator workbook: generated from the manifest, read back after it is filled in."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .manifest import PHOTO_COL

FIELDS_SHEET = "Поля"
FIELDS_HEADER = ["Ключ", "Страница", "№ на схеме", "Где", "Что это (текст в шаблоне)", "Новое значение"]
NO_SLOT_FILL = PatternFill("solid", fgColor="FFD9D9D9")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def display_text(text):
    return (text or "").replace("\r", "\n").replace("\x03", "\n")


def group_columns(group):
    return ["№"] + group["columns"] + ([PHOTO_COL] if group["hasPhotos"] else [])


def instruction_lines(template):
    return [
        f"Шаблон: {template}",
        "",
        "1. Лист «Поля» — разовые тексты альбома (класс, год, школа, учителя, цитаты, обложка).",
        "   Заполняйте колонку «Новое значение». Пустая ячейка = оставить текст из шаблона.",
        "   «№ на схеме» (П1, П2…) — синяя метка на картинке страницы в папке «" + template + "_схема».",
        "2. Листы «Стр …» — повторяющиеся ячейки страницы (фото и подписи под ними).",
        "   Номер строки = красный номер ячейки на схеме этой страницы.",
        "   «Фото (файл)» — имя файла из папки с фото, например ivanova.jpg (расширение можно не писать).",
        "   Серые клетки — у этой ячейки нет такого поля, их можно не заполнять.",
        "   Пустая строка = ячейка останется как в шаблоне.",
        "3. Перенос строки внутри клетки: Alt+Enter.",
        "4. Если текст в шаблоне набран ЗАГЛАВНЫМИ, ваш текст тоже будет переведён в заглавные.",
    ]


def write_template(manifest, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Инструкция"
    for i, line in enumerate(instruction_lines(manifest["template"]), 1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 110

    wf = wb.create_sheet(FIELDS_SHEET)
    wf.append(FIELDS_HEADER)
    for page in manifest["pages"]:
        for f in page["fields"]:
            wf.append([f["key"], page["page"], f"П{f['number']}", f["where"], display_text(f["default"]), None])
    for cell in wf[1]:
        cell.font = BOLD
    for col, width in zip("ABCDEF", (10, 10, 11, 26, 60, 60)):
        wf.column_dimensions[col].width = width
    for row in wf.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP

    for page in manifest["pages"]:
        for g in page["groups"]:
            gs = wb.create_sheet(g["sheet"])
            cols = group_columns(g)
            gs.append(cols)
            for cell in gs[1]:
                cell.font = BOLD
            for n, rec in enumerate(g["records"], 1):
                gs.append([n] + [None] * (len(cols) - 1))
                present = {t["column"] for t in rec["texts"]}
                for ci, col in enumerate(cols[1:], start=2):
                    has = rec["photo"] is not None if col == PHOTO_COL else col in present
                    cell = gs.cell(row=n + 1, column=ci)
                    cell.alignment = WRAP
                    if not has:
                        cell.fill = NO_SLOT_FILL
            gs.column_dimensions["A"].width = 6
            for ci in range(2, len(cols) + 1):
                gs.column_dimensions[gs.cell(row=1, column=ci).column_letter].width = 34
    wb.save(path)


def _cell_str(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value)
    return s if s.strip() else None


def read_filled(manifest, path):
    """{"fields": {key: str}, "groups": {sheet: [{column: str|None}] per record}, "errors", "warnings"}"""
    wb = load_workbook(path, data_only=True)
    data = {"fields": {}, "groups": {}, "errors": [], "warnings": []}

    if FIELDS_SHEET in wb.sheetnames:
        ws = wb[FIELDS_SHEET]
        header = [c.value for c in ws[1]]
        ki, vi = header.index("Ключ"), header.index("Новое значение")
        for row in ws.iter_rows(min_row=2, values_only=True):
            key = row[ki] if ki < len(row) else None
            val = _cell_str(row[vi]) if vi < len(row) else None
            if key and val is not None:
                data["fields"][str(key)] = val
    else:
        data["warnings"].append(f"Нет листа «{FIELDS_SHEET}» — разовые поля останутся как в шаблоне")

    for page in manifest["pages"]:
        for g in page["groups"]:
            sheet = g["sheet"]
            if sheet not in wb.sheetnames:
                data["warnings"].append(f"Нет листа «{sheet}» — ячейки страницы {page['page']} без изменений")
                data["groups"][sheet] = []
                continue
            ws = wb[sheet]
            header = [c.value for c in ws[1]]
            expected = group_columns(g)
            unknown = [h for h in header if h and h not in expected]
            if unknown:
                data["warnings"].append(f"Лист «{sheet}»: неизвестные колонки {unknown} — игнорируются")
            rows = []
            for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                values = {h: _cell_str(v) for h, v in zip(header, row) if h in expected and h != "№"}
                if len(rows) >= len(g["records"]):
                    if any(v is not None for v in values.values()):
                        data["errors"].append(
                            f"Лист «{sheet}», строка {r_idx}: в шаблоне только {len(g['records'])} ячеек")
                    continue
                rows.append(values)
            data["groups"][sheet] = rows
    return data
