"""Synthetic test data: portrait/landscape photos with a visible 'face' mark near the top and a
filled copy of the template Excel (about 80% of cells, some long names, every short field)."""
import argparse
import random
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from autofill.excel_io import FIELDS_SHEET
from autofill.manifest import PHOTO_COL

HERE = Path(__file__).resolve().parent
FIRST = ["Мария", "Иван", "Анна", "Пётр", "Софья", "Артём", "Дарья", "Михаил", "Алиса", "Кирилл"]
LAST = ["Иванова", "Петров", "Смирнова", "Кузнецов", "Попова", "Соколов", "Лебедева", "Константинопольский"]


def make_photo(path, text, portrait, rnd):
    w, h = (1200, 1600) if portrait else (1600, 1100)
    im = Image.new("RGB", (w, h), tuple(rnd.randint(60, 200) for _ in range(3)))
    d = ImageDraw.Draw(im)
    d.ellipse([w / 2 - 180, h * 0.18, w / 2 + 180, h * 0.18 + 360], fill=(240, 220, 200), outline=(0, 0, 0), width=8)
    font = ImageFont.truetype("arialbd.ttf", 90)
    d.text((w / 2, h * 0.75), text, font=font, fill=(255, 255, 255), anchor="mm")
    d.rectangle([0, 0, w - 1, h - 1], outline=(255, 255, 0), width=12)
    im.save(path, quality=90)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--template", required=True)
    ap.add_argument("--fill", type=float, default=0.8, help="доля заполненных ячеек")
    args = ap.parse_args()
    rnd = random.Random(1)
    root = HERE / "demo" / args.template
    photos = root / "photos"
    shutil.rmtree(root, ignore_errors=True)
    photos.mkdir(parents=True)
    xlsx = root / "data.xlsx"
    shutil.copy(HERE / "templates" / f"{args.template}.xlsx", xlsx)
    wb = load_workbook(xlsx)

    for row in wb[FIELDS_SHEET].iter_rows(min_row=2):
        if len(row[4].value or "") <= 30:
            row[5].value = "ТЕСТ " + row[2].value
    n = 0
    for ws in wb.worksheets:
        if not ws.title.startswith("Стр"):
            continue
        header = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2))
        for row in rows[: max(1, int(len(rows) * args.fill))]:
            n += 1
            person = f"{rnd.choice(FIRST)} {rnd.choice(LAST)}"
            for cell, col in zip(row, header):
                if cell.fill.fill_type:  # gray = the cell has no such slot
                    continue
                if col == PHOTO_COL:
                    fname = f"p{n:03d}.jpg"
                    make_photo(photos / fname, str(n), rnd.random() < 0.8, rnd)
                    cell.value = fname
                elif col and col.startswith("Подпись"):
                    cell.value = person.replace(" ", "\n") if "/" in col else person
    wb.save(xlsx)
    print(f"{root}: ячеек {n}, фото {len(list(photos.iterdir()))}")


if __name__ == "__main__":
    main()
