"""Bot class + students -> stage-1 ClassData (for plan/preview/build) and the class workbook (xlsx)."""
from openpyxl import load_workbook

from vignette.config import class_photo_labels
from vignette.names import split_name
from vignette.table import CLASS_NAME, CLASS_SHEET, COMMON_SHEET, STUDENTS_SHEET, write_table

READY = ("filled", "previewed", "approved")


def _path(source, num, kinds=("portrait", "common")):
    """Frame number -> file; a number that cannot be found stays as is (the plan reports it)."""
    if not num:
        return None
    for kind in kinds:
        found = source.find(num, kind)
        if found:
            return found
    return str(num)


def ready_students(store, code):
    return store.students(code, READY)


def to_classdata(cls, students, source, grid_size=None):
    people = []
    for i, st in enumerate(students, start=2):
        surname, first = split_name(st["name"] or "")
        people.append({"name": st["name"], "surname": surname, "first": first,
                       "cover": _path(source, st["cover"], ("portrait",)),
                       "portrait": _path(source, st["portrait"], ("portrait",)),
                       "quote": st["quote"], "group": [_path(source, n, ("common",)) for n in st["grp"] or []],
                       "pages": list(st["pages"] or []), "row": i, "sid": st["id"], "crops": st.get("crops") or {}})
    return {"class_name": cls["title"], "fields": dict(cls["fields"] or {}),
            "class_photos": {k: _path(source, v) for k, v in (cls["class_photos"] or {}).items() if v},
            "students": people, "teachers": [],
            "common_photos": [_path(source, n, ("common",)) for n in cls["common"] or []],
            "grid_size": grid_size or cls.get("expected") or 0, "errors": [], "warnings": []}


def write_class_xlsx(cfg, manifest, cls, students, path):
    """The same workbook an operator fills by hand, with frame numbers as photo names."""
    write_table(cfg, manifest, path)
    wb = load_workbook(path)
    ws = wb[CLASS_SHEET]
    photo_labels = set(class_photo_labels(cfg))
    for row in ws.iter_rows(min_row=2):
        label = row[0].value
        if label == CLASS_NAME:
            row[1].value = cls["title"]
        elif label in photo_labels:
            row[1].value = (cls["class_photos"] or {}).get(label)
        elif label in (cls["fields"] or {}):
            row[1].value = cls["fields"][label]
    titles = {p["page"]: p["title"] for p in cfg["pages"]}
    for st in students:
        pages = ", ".join(titles[p] for p in st["pages"] or [] if p in titles) or "нет"
        grp = st["grp"] or []
        wb[STUDENTS_SHEET].append([st["name"], st["cover"], st["portrait"], st["quote"],
                                   ", ".join(n or "-" for n in grp) if any(grp) else None, pages])
    for num in cls["common"] or []:
        wb[COMMON_SHEET].append([num])
    wb.save(path)
    return path
