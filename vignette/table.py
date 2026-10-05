"""The class workbook: one per class, written empty for a template and read back after filling."""
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from autofill.excel_io import display_text

from .config import class_photo_labels, field_slots, manifest_page, records, split_source
from .names import split_name

CLASS_SHEET = "Класс"
STUDENTS_SHEET = "Ученики"
TEACHERS_SHEET = "Учителя"
COMMON_SHEET = "Общие фото"
PAGES_SHEET = "Страницы"
CLASS_NAME = "Название класса"
STUDENTS_HEADER = ["Фамилия Имя", "Фото обложка", "Фото сетка", "Цитата", "Групповые фото", "Доп. страницы"]
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
LIST_SPLIT = re.compile(r"[,;\n]+")
NO_PAGES = ("-", "нет")          # «Доп. страницы»: only the required pages


def group_slot_count(page_cfg, mpage):
    recs = records(mpage)
    explicit = page_cfg.get("photos", {})
    if page_cfg.get("role", "photos") == "photos":
        return sum(1 for i in range(len(recs)) if split_source(explicit.get(str(i), "group"))[0] == "group")
    return sum(1 for src in explicit.values() if split_source(src)[0] == "group")


def instruction_lines(cfg):
    return [
        f"Шаблон: {cfg['template']}",
        "",
        "1. «Класс» — общие надписи альбома и общие фото класса. Пустая ячейка = как в шаблоне.",
        "2. «Ученики» — строка на ученика: «Фамилия Имя», фото для обложки и для сетки класса (имя файла,",
        "   расширение можно не писать), цитата, групповые фото через запятую (по порядку мест в альбоме),",
        "   доп. страницы через запятую (номер или название из листа «Страницы»; пусто = все, «нет» = только обязательные).",
        "3. «Учителя» — фото и подпись (строки подписи через «/» или Alt+Enter).",
        "4. «Общие фото» — групповые фото по умолчанию: ими добиваются места, если ученик выбрал меньше.",
        "5. Перенос строки внутри клетки: Alt+Enter. PDF получит имя из колонки «Фамилия Имя».",
    ]


def write_table(cfg, manifest, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Инструкция"
    for i, line in enumerate(instruction_lines(cfg), 1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 110

    wc = wb.create_sheet(CLASS_SHEET)
    wc.append(["Поле", "Значение", "Подсказка (как в шаблоне)"])
    wc.append([CLASS_NAME, None, "имя папки с PDF, например 11 А"])
    for f in field_slots(cfg, manifest):
        if not f["personal"]:
            wc.append([f["label"], None, f["hint"] or display_text(f["layers"][0][1]["default"])])
    for label in class_photo_labels(cfg):
        wc.append([label, None, "имя файла фото"])
    _style(wc, (44, 40, 60))

    wsd = wb.create_sheet(STUDENTS_SHEET)
    wsd.append(STUDENTS_HEADER)
    _style(wsd, (32, 18, 18, 50, 40, 24))

    if any(p.get("role") == "teachers" for p in cfg["pages"]):
        wt = wb.create_sheet(TEACHERS_SHEET)
        wt.append(["Фото", "Подпись"])
        _style(wt, (24, 60))

    wg = wb.create_sheet(COMMON_SHEET)
    wg.append(["Файл"])
    _style(wg, (40,))

    wp = wb.create_sheet(PAGES_SHEET)
    wp.append(["№", "Название", "Обязательная", "Групповых фото"])
    for p in cfg["pages"]:
        wp.append([p["page"], p["title"], "да" if p.get("required") else "нет",
                   group_slot_count(p, manifest_page(manifest, p["page"]))])
    _style(wp, (10, 32, 14, 16))
    wb.save(path)


def _style(ws, widths):
    for cell in ws[1]:
        cell.font = BOLD
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP


def cell_str(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\r\n", "\n").strip()
    return s or None


def split_list(v):
    s = cell_str(v)
    return [x.strip() for x in LIST_SPLIT.split(s) if x.strip()] if s else []


def page_lookup(cfg):
    table = {}
    for p in cfg["pages"]:
        table[p["page"].casefold()] = p["page"]
        table[p["title"].casefold()] = p["page"]
        if p["page"].isdigit():
            table[str(int(p["page"]))] = p["page"]
    return table


def _rows(wb, sheet):
    if sheet not in wb.sheetnames:
        return []
    return list(wb[sheet].iter_rows(min_row=2, values_only=True))


def read_table(cfg, manifest, path):
    wb = load_workbook(path, data_only=True)
    errors, warnings = [], []
    known = {f["label"] for f in field_slots(cfg, manifest) if not f["personal"]}
    photo_labels = set(class_photo_labels(cfg))
    fields = {label: None for label in known}
    class_photos, class_name = {}, None
    for row in _rows(wb, CLASS_SHEET):
        label, value = cell_str(row[0]) if row else None, cell_str(row[1]) if len(row) > 1 else None
        if label is None:
            continue
        if label == CLASS_NAME:
            class_name = value
        elif label in known:
            fields[label] = value
        elif label in photo_labels:
            class_photos[label] = value
        elif value is not None:
            warnings.append(f"Лист «{CLASS_SHEET}»: неизвестное поле «{label}» — пропущено")

    pages = page_lookup(cfg)
    students, seen = [], {}
    for n, row in enumerate(_rows(wb, STUDENTS_SHEET), start=2):
        row = list(row) + [None] * (6 - len(row))
        name, cover, portrait, quote = (cell_str(v) for v in row[:4])
        if not any(cell_str(v) for v in row[:6]):
            continue
        if name is None:
            errors.append(f"Лист «{STUDENTS_SHEET}», строка {n}: не указано имя")
            continue
        chosen = None
        if split_list(row[5]):
            chosen = []
            for token in split_list(row[5]):
                if token.casefold() in NO_PAGES:
                    continue
                page = pages.get(token.casefold())
                if page is None:
                    errors.append(f"Лист «{STUDENTS_SHEET}», строка {n}: нет страницы «{token}»")
                elif page not in chosen:
                    chosen.append(page)
        name = " ".join(name.split())
        if name in seen:
            seen[name] += 1
            warnings.append(f"Ученик «{name}» записан дважды — второй получит имя «{name} ({seen[name]})»")
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        surname, first = split_name(name.split(" (")[0])
        students.append({"name": name, "surname": surname, "first": first, "cover": cover, "portrait": portrait,
                         "quote": quote, "group": split_list(row[4]), "pages": chosen, "row": n})

    teachers = []
    for row in _rows(wb, TEACHERS_SHEET):
        photo, caption = (cell_str(v) for v in (list(row) + [None, None])[:2])
        if photo or caption:
            caption = "\n".join(x.strip() for x in re.split(r"[/\n]", caption or "") if x.strip())
            teachers.append({"photo": photo, "caption": caption})
    common = [cell_str(r[0]) for r in _rows(wb, COMMON_SHEET) if r and cell_str(r[0])]
    return {"class_name": class_name or Path(path).stem, "fields": fields, "class_photos": class_photos,
            "students": students, "teachers": teachers, "common_photos": common,
            "errors": errors, "warnings": warnings}
